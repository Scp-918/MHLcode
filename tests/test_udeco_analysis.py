from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
import pytest
import matplotlib
import nbformat
from openpyxl import load_workbook
from scipy.io import savemat

matplotlib.use("Agg")

from udeco_analysis import (
    AnalysisConfig,
    PYDISPLAY_SENSOR_COLUMNS,
    analyze_all_subjects,
    analyze_sample,
    analyze_subject,
    classify_udeco,
    detect_motion_stage,
    discover_subject_samples,
    export_excel,
    load_calibration,
    load_udeco_sample,
    plot_validation,
    robust_sigma,
    select_validation_candidates,
    summarize_subject,
)


def write_calibration(root: Path, *, c2: float = 2.0, c3: float = 3.0) -> Path:
    c_result = np.empty((3, 1), dtype=object)
    c_result[:, 0] = [
        {"channel": 1, "c": 1.0},
        {"channel": 2, "c": c2},
        {"channel": 3, "c": c3},
    ]
    k_result = np.empty((2, 1), dtype=object)
    k_result[:, 0] = [
        {"mainChannel": 2, "k": 0.1},
        {"mainChannel": 3, "k": 0.2},
    ]
    path = root / "ck.mat"
    savemat(path, {"cResult": c_result, "kResult": k_result})
    return path


def write_sensor(
    path: Path,
    *,
    absolute_seq: list[int] | None = None,
    parser_valid: list[bool] | None = None,
    n: int = 4,
) -> None:
    absolute_seq = absolute_seq or list(range(n))
    parser_valid = parser_valid or [True] * len(absolute_seq)
    rows = len(absolute_seq)
    frame = pd.DataFrame({name: np.ones(rows) for name in PYDISPLAY_SENSOR_COLUMNS})
    frame["frame_seq"] = np.asarray(absolute_seq) % 65536
    frame["absolute_seq_u64"] = absolute_seq
    frame["segment_id"] = 0
    frame["sample_seq"] = np.arange(rows) + 1
    frame["parser_valid"] = parser_valid
    frame["Uh2"] = 4.0
    frame["Uh3"] = 6.0
    frame["Uc2"] = 0.5
    frame["Uc3"] = 1.0
    frame["ACC_X"] = 0.0
    frame["ACC_Y"] = 0.0
    frame["ACC_Z"] = 1.0
    frame.to_csv(path, index=False)


def write_hr(path: Path) -> None:
    pd.DataFrame({"elapsed_seconds": [0.0, 1.0], "hr_bpm": [70, 71]}).to_csv(path, index=False)


def test_load_calibration_and_udeco_formula_with_gap_repair(tmp_path: Path) -> None:
    subject_dir = tmp_path / "ts"
    subject_dir.mkdir()
    write_calibration(tmp_path)
    sensor = subject_dir / "ts_write_1_sensor.csv"
    write_sensor(sensor, absolute_seq=[0, 2, 3], parser_valid=[True, True, False])

    calibration = load_calibration(tmp_path)
    sample = load_udeco_sample(sensor, data_root=tmp_path, fs=100)

    assert calibration.c2 == pytest.approx(2.0)
    assert calibration.c3 == pytest.approx(3.0)
    np.testing.assert_allclose(sample.time_s, [0.0, 0.01, 0.02])
    np.testing.assert_allclose(sample.udeco2, (4.0 - 0.5) / (2.0 - 0.5))
    np.testing.assert_allclose(sample.udeco3, (6.0 - 1.0) / (3.0 - 1.0))
    assert int(sample.qc["InterpFlag"].sum()) == 1
    assert sample.metadata["removed_parser_invalid_rows"] == 1
    assert sample.metadata["inserted_missing_rows"] == 1


def test_near_zero_udeco_denominator_is_interpolated_and_flagged(tmp_path: Path) -> None:
    subject_dir = tmp_path / "ts"
    subject_dir.mkdir()
    write_calibration(tmp_path)
    sensor = subject_dir / "ts_write_1_sensor.csv"
    write_sensor(sensor, n=3)
    frame = pd.read_csv(sensor)
    frame.loc[1, "Uc2"] = 2.0
    frame.to_csv(sensor, index=False)

    sample = load_udeco_sample(sensor, data_root=tmp_path)

    assert np.isfinite(sample.udeco2).all()
    assert sample.qc["ud_denominator_near_zero"].tolist() == [0, 1, 0]


def test_discovery_requires_strict_sensor_hr_pairs(tmp_path: Path) -> None:
    subject_dir = tmp_path / "ts"
    subject_dir.mkdir()
    write_sensor(subject_dir / "ts_write_1_sensor.csv")
    write_hr(subject_dir / "ts_write_1_HRdata.csv")
    write_sensor(subject_dir / "ts_run_2_sensor.csv")
    write_sensor(subject_dir / "other_rope_1_sensor.csv")
    (subject_dir / "notes.csv").write_text("a,b\n1,2\n", encoding="utf-8")

    pairs, issues = discover_subject_samples(tmp_path, "ts")

    assert [pair.sample_name for pair in pairs] == ["ts_write_1"]
    categories = {issue.file_name: issue.category for issue in issues}
    assert categories["ts_run_2_sensor.csv"] == "missing_pair"
    assert categories["other_rope_1_sensor.csv"] == "subject_mismatch"
    assert categories["notes.csv"] == "invalid_name"


def test_missing_calibration_is_explicit(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="ck.mat"):
        load_calibration(tmp_path)


def test_robust_sigma_uses_epsilon_for_constant_signal() -> None:
    assert robust_sigma(np.ones(20), epsilon=1e-7) == pytest.approx(1e-7)


def test_classify_udeco_preserves_100_ms_h_state_and_nan_for_missing_l() -> None:
    fs = 100
    time = np.arange(500, dtype=float) / fs
    rng = np.random.default_rng(7)
    raw = rng.normal(0.0, 0.005, time.size)
    short_h = (time >= 2.0) & (time < 2.1)
    raw[short_h] += 0.8
    config = AnalysisConfig(initial_m_sec=1.0)

    result = classify_udeco(raw, time=time, fs=fs, config=config)

    assert np.mean(result["labels"][short_h] == "H") >= 0.8
    assert result["statistics"]["n_detected_states"] == 2
    assert np.isnan(result["statistics"]["L_mean"])
    assert np.isnan(result["statistics"]["L_total_duration"])
    assert result["statistics"]["H_total_duration"] == pytest.approx(0.1)


def test_classify_udeco_preserves_200_ms_l_state() -> None:
    fs = 100
    time = np.arange(500, dtype=float) / fs
    raw = np.zeros(time.size)
    short_l = (time >= 2.0) & (time < 2.2)
    raw[short_l] = -0.8

    result = classify_udeco(raw, time=time, fs=fs, config=AnalysisConfig(initial_m_sec=1.0))

    assert np.mean(result["labels"][short_l] == "L") >= 0.9
    assert result["statistics"]["L_total_duration"] == pytest.approx(0.2)


@pytest.mark.filterwarnings("error")
def test_optional_gmm_path_runs_on_segment_medians_only() -> None:
    fs = 100
    time = np.arange(500, dtype=float) / fs
    raw = np.zeros(time.size)
    raw[(time >= 2.0) & (time < 2.5)] = 1.0

    result = classify_udeco(
        raw,
        time=time,
        fs=fs,
        config=AnalysisConfig(initial_m_sec=1.0, use_gmm=True),
    )

    assert {"M", "H"} == set(result["labels"])
    assert len(result["segments"]) < len(raw)


def test_classify_udeco_rejects_single_sample_spike() -> None:
    fs = 100
    time = np.arange(500, dtype=float) / fs
    raw = np.zeros(time.size)
    raw[250] = 5.0

    result = classify_udeco(raw, time=time, fs=fs, config=AnalysisConfig(initial_m_sec=1.0))

    assert set(result["labels"]) == {"M"}
    assert result["statistics"]["n_detected_states"] == 1


def test_dynamic_m_tracks_slow_drift_without_false_h_or_l() -> None:
    fs = 100
    time = np.arange(2000, dtype=float) / fs
    rng = np.random.default_rng(9)
    raw = 0.01 * time + rng.normal(0.0, 0.006, time.size)
    config = AnalysisConfig(initial_m_sec=2.0, m_baseline_smooth_sec=5.0)

    result = classify_udeco(raw, time=time, fs=fs, config=config)

    assert set(result["labels"]) == {"M"}
    assert result["M_t"][-1] - result["M_t"][0] > 0.15


def test_three_states_and_direct_h_l_transition_are_retained_and_flagged() -> None:
    fs = 100
    time = np.arange(500, dtype=float) / fs
    raw = np.zeros(time.size)
    raw[(time >= 2.0) & (time < 2.5)] = 1.0
    raw[(time >= 2.5) & (time < 3.0)] = -1.0
    config = AnalysisConfig(initial_m_sec=1.0)

    result = classify_udeco(raw, time=time, fs=fs, config=config)

    assert {"H", "M", "L"} == set(result["labels"])
    assert "direct_H_to_L" in result["statistics"]["QC_flag"]
    assert result["statistics"]["transition_count"] >= 3
    assert len(result["change_points"]) >= 3


def test_motion_detection_finds_rest_motion_recovery() -> None:
    fs = 100
    time = np.arange(3000, dtype=float) / fs
    rng = np.random.default_rng(12)
    accx = rng.normal(0.0, 0.01, time.size)
    accy = rng.normal(0.0, 0.01, time.size)
    accz = rng.normal(0.0, 0.01, time.size)
    motion = (time >= 10.0) & (time < 20.0)
    accx[motion] += np.sin(2 * np.pi * 2.0 * time[motion])
    config = AnalysisConfig(
        motion_calibration_sec=5.0,
        motion_window_sec=1.0,
        motion_transition_windows=2,
    )

    result = detect_motion_stage(accx, accy, accz, fs=fs, config=config)

    assert result["status"] == "ok"
    assert result["motion_start_s"] == pytest.approx(10.0, abs=1.0)
    assert result["motion_end_s"] == pytest.approx(20.0, abs=1.0)


def test_analyze_sample_keeps_udeco_channels_independent_and_plots(tmp_path: Path) -> None:
    subject_dir = tmp_path / "ts"
    subject_dir.mkdir()
    write_calibration(tmp_path)
    sensor = subject_dir / "ts_write_1_sensor.csv"
    write_sensor(sensor, n=500)
    frame = pd.read_csv(sensor)
    frame.loc[200:209, "Uh2"] += 1.5
    frame.loc[300:319, "Uh3"] -= 2.0
    frame.to_csv(sensor, index=False)

    result = analyze_sample(sensor, data_root=tmp_path, config=AnalysisConfig(initial_m_sec=1.0))
    main_fig, debug_fig = plot_validation(result, show_debug_info=True)

    assert set(result["channels"]) == {"Udeco2", "Udeco3"}
    assert "H" in result["channels"]["Udeco2"]["labels"]
    assert "L" not in result["channels"]["Udeco2"]["labels"]
    assert "L" in result["channels"]["Udeco3"]["labels"]
    assert "H" not in result["channels"]["Udeco3"]["labels"]
    assert len(main_fig.axes) == 2
    assert len(debug_fig.axes) == 8


def test_analyze_subject_keeps_bad_sample_as_error_rows(tmp_path: Path) -> None:
    subject_dir = tmp_path / "ts"
    subject_dir.mkdir()
    write_calibration(tmp_path)
    good_sensor = subject_dir / "ts_write_1_sensor.csv"
    write_sensor(good_sensor, n=500)
    write_hr(subject_dir / "ts_write_1_HRdata.csv")
    bad_sensor = subject_dir / "ts_run_2_sensor.csv"
    pd.DataFrame({"wrong": [1, 2]}).to_csv(bad_sensor, index=False)
    write_hr(subject_dir / "ts_run_2_HRdata.csv")

    result = analyze_subject(tmp_path, "ts", config=AnalysisConfig(initial_m_sec=1.0))

    rows = result["sample_results"]
    assert len(rows) == 4
    assert set(rows.loc[rows["sample_name"] == "ts_write_1", "analysis_status"]) == {"ok"}
    bad = rows.loc[rows["sample_name"] == "ts_run_2"]
    assert set(bad["channel"]) == {"Udeco2", "Udeco3"}
    assert set(bad["analysis_status"]) == {"error"}
    assert bad["H_mean"].isna().all()
    assert "analysis_error" in set(result["qc"]["category"])


def test_analyze_all_subjects_keeps_missing_subject_as_subject_qc(tmp_path: Path) -> None:
    batch = analyze_all_subjects(tmp_path, ["ts", "lyx"])

    assert set(batch) == {"ts", "lyx"}
    assert batch["ts"]["sample_results"].empty
    assert batch["lyx"]["sample_results"].empty
    assert set(batch["ts"]["qc"]["category"]) == {"subject_error"}


def test_subject_summary_ignores_nan_and_counts_valid_samples() -> None:
    sample_results = pd.DataFrame(
        [
            {"sample_name": "a", "subject": "ts", "channel": "Udeco2", "analysis_status": "ok", "H_mean": 1.0, "M_mean": 0.0, "L_mean": np.nan},
            {"sample_name": "b", "subject": "ts", "channel": "Udeco2", "analysis_status": "ok", "H_mean": 3.0, "M_mean": 0.2, "L_mean": np.nan},
            {"sample_name": "c", "subject": "ts", "channel": "Udeco2", "analysis_status": "error", "H_mean": 99.0, "M_mean": 99.0, "L_mean": -99.0},
            {"sample_name": "a", "subject": "ts", "channel": "Udeco3", "analysis_status": "ok", "H_mean": np.nan, "M_mean": 2.0, "L_mean": -1.0},
        ]
    )

    summary = summarize_subject(sample_results)
    h2 = summary.loc[(summary["channel"] == "Udeco2") & (summary["state"] == "H")].iloc[0]
    l2 = summary.loc[(summary["channel"] == "Udeco2") & (summary["state"] == "L")].iloc[0]

    assert h2["mean"] == pytest.approx(2.0)
    assert h2["std"] == pytest.approx(np.sqrt(2.0))
    assert h2["N"] == 2
    assert l2["N"] == 0
    assert np.isnan(l2["mean"])


def test_export_excel_writes_six_subject_sheets_and_three_parts(tmp_path: Path) -> None:
    subjects = ["ts", "lyx", "lzj", "cgx", "pjy", "hjy"]
    rows = pd.DataFrame(
        [
            {
                "sample_name": "ts_write_1",
                "subject": "ts",
                "motion_type": "write",
                "motion_index": 1,
                "channel": channel,
                "analysis_status": "ok",
                "H_mean": 1.0 if channel == "Udeco2" else np.nan,
                "M_mean": 0.0,
                "L_mean": np.nan,
                "QC_flag": "OK",
            }
            for channel in ("Udeco2", "Udeco3")
        ]
    )
    batch = {}
    for subject in subjects:
        frame = rows.copy() if subject == "ts" else rows.iloc[0:0].copy()
        batch[subject] = {
            "subject": subject,
            "sample_results": frame,
            "summary": summarize_subject(frame),
            "qc": pd.DataFrame(columns=["file_name", "category", "reason"]),
            "sample_details": {},
        }
    output = tmp_path / "udeco_results.xlsx"

    written = export_excel(batch, output)

    assert written == output
    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == subjects
    values = [cell.value for row in workbook["ts"].iter_rows() for cell in row]
    assert "Part A: Sample-level results" in values
    assert "Part B: Subject summary" in values
    assert "Part C: Skipped files / QC" in values


def test_select_validation_candidates_covers_key_manual_checks() -> None:
    def channel_result(labels: list[str], durations: list[float], *, drift: float = 0.0) -> dict:
        return {
            "labels": np.asarray(labels, dtype=object),
            "M_t": np.asarray([0.0, drift]),
            "statistics": {
                "H_mean": 1.0 if "H" in labels else np.nan,
                "L_mean": -1.0 if "L" in labels else np.nan,
            },
            "segments": pd.DataFrame({"final_label": labels, "duration": durations}),
        }

    batch = {
        "ts": {
            "sample_details": {
                "ts_write_1": {
                    "sample_name": "ts_write_1",
                    "subject": "ts",
                    "metadata": {"sensor_csv": "ts_write_1_sensor.csv"},
                    "motion": {"window_std": np.asarray([0.1, 2.0])},
                    "channels": {
                        "Udeco2": channel_result(["M", "H", "M"], [1.0, 0.1, 1.0], drift=0.4),
                        "Udeco3": channel_result(["M", "L", "M"], [1.0, 0.2, 1.0]),
                    },
                }
            }
        }
    }

    candidates = select_validation_candidates(batch)

    assert {"M_plus_H", "M_plus_L", "shortest_non_M", "strongest_motion", "largest_M_drift"} <= set(
        candidates["category"]
    )


def test_notebook_has_required_configuration_and_workflow_sections() -> None:
    path = Path("udeco_state_analysis.ipynb")
    assert path.is_file()
    notebook = nbformat.read(path, as_version=4)
    nbformat.validate(notebook)
    for cell in notebook.cells:
        if cell.cell_type == "code":
            assert cell.execution_count is None
            assert cell.outputs == []
    markdown = "\n".join(cell.source for cell in notebook.cells if cell.cell_type == "markdown")
    code = "\n".join(cell.source for cell in notebook.cells if cell.cell_type == "code")

    for section in (
        "用户参数区",
        "Validation / Manual Inspection",
        "批量分析",
        "Excel 输出",
        "汇总检查",
    ):
        assert section in markdown
    for setting in (
        "DATA_ROOT",
        "SUBJECTS",
        "MEDIAN_FILTER_SEC",
        "CHANGE_WINDOW_SEC",
        "ENTER_SIGMA",
        "EXIT_SIGMA",
        "USE_GMM",
        "VALIDATION_FILE",
        "SHOW_DEBUG_INFO",
        "RUN_BATCH",
        "OUTPUT_EXCEL",
    ):
        assert setting in code
    assert "analyze_sample" in code
    assert "analyze_all_subjects" in code
    assert "export_excel" in code
