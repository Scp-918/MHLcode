from __future__ import annotations

from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from scipy.io import savemat

matplotlib.use("Agg")

from udeco_analysis import PYDISPLAY_SENSOR_COLUMNS
from udeco_interval_analysis import (
    IntervalAnalysisConfig,
    analyze_interval_batch,
    analyze_interval_sample,
    compute_channel_statistics,
    export_interval_excel,
    export_interval_means_excel,
    plot_interval_validation,
)


def _write_calibration(root: Path) -> None:
    c_result = np.empty((3, 1), dtype=object)
    c_result[:, 0] = [
        {"channel": 1, "c": 1.0},
        {"channel": 2, "c": 2.0},
        {"channel": 3, "c": 3.0},
    ]
    k_result = np.empty((2, 1), dtype=object)
    k_result[:, 0] = [
        {"mainChannel": 2, "k": 0.1},
        {"mainChannel": 3, "k": 0.2},
    ]
    savemat(root / "ck.mat", {"cResult": c_result, "kResult": k_result})


def _write_hr(path: Path) -> None:
    pd.DataFrame({"elapsed_seconds": [0.0, 1.0], "hr_bpm": [70, 71]}).to_csv(path, index=False)


def _write_sensor_with_udeco(path: Path, udeco2: np.ndarray, udeco3: np.ndarray) -> None:
    assert udeco2.shape == udeco3.shape
    n = udeco2.size
    frame = pd.DataFrame({name: np.ones(n) for name in PYDISPLAY_SENSOR_COLUMNS})
    frame["frame_seq"] = np.arange(n) % 65536
    frame["absolute_seq_u64"] = np.arange(n)
    frame["segment_id"] = 0
    frame["sample_seq"] = np.arange(n) + 1
    frame["parser_valid"] = True
    frame["Uc2"] = 0.5
    frame["Uc3"] = 1.0
    frame["Uh2"] = 0.5 + 1.5 * udeco2
    frame["Uh3"] = 1.0 + 2.0 * udeco3
    frame["ACC_X"] = 0.0
    frame["ACC_Y"] = 0.0
    frame["ACC_Z"] = 1.0
    frame.to_csv(path, index=False)


def test_compute_channel_statistics_uses_half_open_windows_and_sample_variance() -> None:
    fs = 1
    time = np.arange(50, dtype=float)
    signal = np.zeros(50, dtype=float)
    signal[15:35] = np.arange(20, dtype=float)
    signal[35] = 1000.0  # 35 s 必须排除在运动统计区间之外
    config = IntervalAnalysisConfig(median_filter_sec=1.0)

    result = compute_channel_statistics(signal, time, motion_start_s=0.0, fs=fs, config=config)
    stats = result["statistics"]

    assert stats["pre30_mean"] == pytest.approx(3.5)
    assert stats["pre30_var"] == pytest.approx(22.32758620689655)
    assert stats["interval_mean"] == pytest.approx(9.5)
    assert stats["interval_max"] == pytest.approx(19.0)
    assert stats["interval_min"] == pytest.approx(0.0)
    assert stats["threshold1"] == pytest.approx(14.25)
    assert stats["threshold2"] == pytest.approx(6.65)
    assert stats["upper_mean"] == pytest.approx(17.0)
    assert stats["upper_var"] == pytest.approx(2.5)
    assert stats["lower_mean"] == pytest.approx(3.0)
    assert stats["lower_var"] == pytest.approx(4.666666666666667)
    assert stats["upper_n"] == 5
    assert stats["lower_n"] == 7
    assert not result["masks"]["interval"][35]


def test_constant_signal_reports_empty_strict_threshold_groups() -> None:
    fs = 10
    time = np.arange(500, dtype=float) / fs
    signal = np.full(time.size, 2.5)

    result = compute_channel_statistics(
        signal,
        time,
        motion_start_s=0.0,
        fs=fs,
        config=IntervalAnalysisConfig(median_filter_sec=0.1),
    )
    stats = result["statistics"]

    assert stats["threshold1"] == pytest.approx(2.5)
    assert stats["threshold2"] == pytest.approx(2.5)
    assert stats["upper_n"] == 0
    assert stats["lower_n"] == 0
    assert np.isnan(stats["upper_mean"])
    assert np.isnan(stats["upper_var"])
    assert np.isnan(stats["lower_mean"])
    assert np.isnan(stats["lower_var"])
    assert "upper_no_points" in result["qc_flags"]
    assert "lower_no_points" in result["qc_flags"]


def test_short_record_does_not_silently_truncate_motion_window() -> None:
    fs = 10
    time = np.arange(300, dtype=float) / fs
    signal = np.linspace(0.0, 1.0, time.size)

    result = compute_channel_statistics(
        signal,
        time,
        motion_start_s=10.0,
        fs=fs,
        config=IntervalAnalysisConfig(median_filter_sec=0.1),
    )

    assert "motion_interval_out_of_range" in result["qc_flags"]
    assert np.isnan(result["statistics"]["threshold1"])
    assert np.isnan(result["statistics"]["upper_mean"])


def test_sample_falls_back_to_70_seconds_and_keeps_channels_independent(tmp_path: Path) -> None:
    fs = 10
    n = 110 * fs
    time = np.arange(n, dtype=float) / fs
    udeco2 = np.ones(n)
    udeco3 = np.full(n, 10.0)
    interval = (time >= 85.0) & (time < 105.0)
    udeco2[interval] = np.linspace(0.0, 2.0, interval.sum())
    udeco3[interval] = np.linspace(8.0, 14.0, interval.sum())
    subject_dir = tmp_path / "ts"
    subject_dir.mkdir()
    _write_calibration(tmp_path)
    sensor = subject_dir / "ts_write_1_sensor.csv"
    _write_sensor_with_udeco(sensor, udeco2, udeco3)

    result = analyze_interval_sample(
        sensor,
        data_root=tmp_path,
        fs=fs,
        config=IntervalAnalysisConfig(median_filter_sec=0.1),
    )

    assert result["motion_start_s"] == pytest.approx(70.0)
    assert result["motion_start_source"] == "fallback_70s"
    assert "motion_detection_failed_fallback_70s" in result["QC_flag"]
    assert result["channels"]["Udeco2"]["statistics"]["interval_mean"] == pytest.approx(1.0)
    assert result["channels"]["Udeco3"]["statistics"]["interval_mean"] == pytest.approx(11.0)


def test_plot_and_excel_expose_regions_and_fallback_marker(tmp_path: Path) -> None:
    fs = 10
    n = 110 * fs
    time = np.arange(n, dtype=float) / fs
    subject_dir = tmp_path / "ts"
    subject_dir.mkdir()
    _write_calibration(tmp_path)
    sensor = subject_dir / "ts_write_1_sensor.csv"
    hr = subject_dir / "ts_write_1_HRdata.csv"
    _write_sensor_with_udeco(sensor, 1.0 + np.sin(time), 2.0 + np.cos(time))
    _write_hr(hr)
    config = IntervalAnalysisConfig(median_filter_sec=0.1)

    sample = analyze_interval_sample(sensor, data_root=tmp_path, fs=fs, config=config)
    figure = plot_interval_validation(sample)
    batch = analyze_interval_batch(tmp_path, ["ts", "lyx"], fs=fs, config=config)
    output = export_interval_excel(batch, tmp_path / "result.xlsx")

    assert len(figure.axes) == 2
    assert all(len(axis.patches) >= 2 for axis in figure.axes)
    assert all(len(axis.collections) >= 2 for axis in figure.axes)
    assert output.is_file()
    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["ts", "lyx"]
    ts_sheet = workbook["ts"]
    headers = [cell.value for cell in ts_sheet[3]]
    assert "Udeco2_pre30_mean" in headers
    assert "Udeco3_lower_var" in headers
    source_col = headers.index("motion_start_source") + 1
    assert ts_sheet.cell(4, source_col).value == "fallback_70s"


def test_compact_excel_contains_chinese_motion_type_and_six_means(tmp_path: Path) -> None:
    columns = [
        "运动类型",
        "Udeco2_pre30_mean",
        "Udeco2_upper_mean",
        "Udeco2_lower_mean",
        "Udeco3_pre30_mean",
        "Udeco3_upper_mean",
        "Udeco3_lower_mean",
    ]
    sample_results = pd.DataFrame(
        [
            {
                "sample_name": f"ts_{motion_type}_1",
                "motion_type": motion_type,
                "Udeco2_pre30_mean": 1.0,
                "Udeco2_upper_mean": 2.0,
                "Udeco2_lower_mean": 0.5,
                "Udeco3_pre30_mean": 10.0,
                "Udeco3_upper_mean": 12.0,
                "Udeco3_lower_mean": 8.0,
                "Udeco2_pre30_var": 99.0,
                "QC_flag": "must_not_be_exported",
            }
            for motion_type in ("write", "gripper", "run", "rope")
        ]
    )
    batch = {
        "ts": {"sample_results": sample_results},
        "lyx": {"sample_results": pd.DataFrame(columns=sample_results.columns)},
    }

    output = export_interval_means_excel(batch, tmp_path / "means_only.xlsx")

    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["ts", "lyx"]
    assert [cell.value for cell in workbook["ts"][1]] == columns
    assert [workbook["ts"].cell(row, 1).value for row in range(2, 6)] == ["写字", "抓握", "跑步", "跳绳"]
    assert [cell.value for cell in workbook["ts"][2]][1:] == [1.0, 2.0, 0.5, 10.0, 12.0, 8.0]
    assert [cell.value for cell in workbook["lyx"][1]] == columns
