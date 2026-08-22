"""Udeco2/Udeco3 固定时间区间统计、验证绘图与 Excel 导出。

本模块只复用 :mod:`udeco_analysis` 的数据读取、文件发现和 ACC 运动检测，
不调用 H/M/L 状态分类流程。统计值均来自轻度中值滤波后的 Udeco 信号。
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

import numpy as np
import pandas as pd
from scipy.ndimage import median_filter

from udeco_analysis import (
    AnalysisConfig,
    detect_motion_stage,
    discover_subject_samples,
    load_udeco_sample,
)


CORE_STATISTICS = (
    "pre30_mean",
    "pre30_var",
    "upper_mean",
    "upper_var",
    "lower_mean",
    "lower_var",
)

DIAGNOSTIC_STATISTICS = (
    "interval_mean",
    "interval_max",
    "interval_min",
    "threshold1",
    "threshold2",
    "upper_n",
    "lower_n",
)

COMPACT_MEAN_COLUMNS = (
    "motion_type",
    "Udeco2_pre30_mean",
    "Udeco2_upper_mean",
    "Udeco2_lower_mean",
    "Udeco3_pre30_mean",
    "Udeco3_upper_mean",
    "Udeco3_lower_mean",
)

MOTION_TYPE_LABELS = {
    "write": "写字",
    "gripper": "抓握",
    "run": "跑步",
    "rope": "跳绳",
}


@dataclass(frozen=True)
class IntervalAnalysisConfig:
    """固定区间统计参数；所有时间参数均使用秒。"""

    median_filter_sec: float = 0.03
    pre_duration_sec: float = 30.0
    motion_offset_start_sec: float = 15.0
    motion_offset_end_sec: float = 35.0
    fallback_motion_start_sec: float = 70.0
    epsilon: float = 1e-9
    motion_calibration_sec: float = 30.0
    motion_window_sec: float = 8.0
    motion_step_sec: float = 1.0
    motion_threshold_sigma: float = 3.0
    motion_transition_windows: int = 6

    def __post_init__(self) -> None:
        if self.median_filter_sec < 0:
            raise ValueError("median_filter_sec must be non-negative")
        if self.pre_duration_sec <= 0:
            raise ValueError("pre_duration_sec must be positive")
        if self.motion_offset_start_sec < 0:
            raise ValueError("motion_offset_start_sec must be non-negative")
        if self.motion_offset_end_sec <= self.motion_offset_start_sec:
            raise ValueError("motion_offset_end_sec must be greater than motion_offset_start_sec")
        if self.fallback_motion_start_sec < 0:
            raise ValueError("fallback_motion_start_sec must be non-negative")


def compute_channel_statistics(
    signal: np.ndarray,
    time_s: np.ndarray,
    motion_start_s: float,
    fs: int,
    config: IntervalAnalysisConfig | None = None,
) -> dict[str, Any]:
    """计算单通道前 30 s、运动区间上侧和下侧的均值及样本方差。

    时间区间采用半开区间：``[0, 30)`` 与
    ``[motion_start+15, motion_start+35)``。阈值相等的点不进入上、下组。
    """

    config = config or IntervalAnalysisConfig()
    raw = np.asarray(signal, dtype=float)
    time = np.asarray(time_s, dtype=float)
    if raw.ndim != 1 or time.ndim != 1 or raw.size != time.size or raw.size == 0:
        raise ValueError("signal and time_s must be equal-length, non-empty 1D arrays")
    if int(fs) <= 0:
        raise ValueError("fs must be positive")
    if not np.isfinite(time).all() or np.any(np.diff(time) <= 0):
        raise ValueError("time_s must be finite and strictly increasing")

    qc_flags: list[str] = []
    finite = np.isfinite(raw)
    if not finite.any():
        raise ValueError("signal has no finite values")
    if not finite.all():
        raw = pd.Series(raw).interpolate(limit_direction="both").to_numpy(dtype=float)
        qc_flags.append("signal_nonfinite_interpolated")

    filter_samples = _seconds_to_odd_samples(config.median_filter_sec, int(fs))
    filtered = median_filter(raw, size=filter_samples, mode="nearest") if filter_samples > 1 else raw.copy()

    pre_start = 0.0
    pre_end = float(config.pre_duration_sec)
    interval_start = float(motion_start_s + config.motion_offset_start_sec)
    interval_end = float(motion_start_s + config.motion_offset_end_sec)
    pre_mask = (time >= pre_start) & (time < pre_end)
    interval_mask = (time >= interval_start) & (time < interval_end)
    upper_mask = np.zeros(raw.size, dtype=bool)
    lower_mask = np.zeros(raw.size, dtype=bool)

    statistics: dict[str, float | int] = {
        "pre30_mean": float("nan"),
        "pre30_var": float("nan"),
        "interval_mean": float("nan"),
        "interval_max": float("nan"),
        "interval_min": float("nan"),
        "threshold1": float("nan"),
        "threshold2": float("nan"),
        "upper_mean": float("nan"),
        "upper_var": float("nan"),
        "lower_mean": float("nan"),
        "lower_var": float("nan"),
        "upper_n": 0,
        "lower_n": 0,
    }

    if _covers_half_open_interval(time, pre_start, pre_end, int(fs), config.epsilon):
        statistics["pre30_mean"], statistics["pre30_var"] = _mean_and_sample_variance(filtered[pre_mask])
    else:
        qc_flags.append("pre30_interval_out_of_range")

    if _covers_half_open_interval(time, interval_start, interval_end, int(fs), config.epsilon):
        interval_values = filtered[interval_mask]
        interval_mean = float(np.mean(interval_values))
        interval_max = float(np.max(interval_values))
        interval_min = float(np.min(interval_values))
        threshold1 = 0.5 * interval_mean + 0.5 * interval_max
        threshold2 = 0.7 * interval_mean + 0.3 * interval_min
        upper_mask = interval_mask & (filtered > threshold1)
        lower_mask = interval_mask & (filtered < threshold2)
        upper_mean, upper_var = _mean_and_sample_variance(filtered[upper_mask])
        lower_mean, lower_var = _mean_and_sample_variance(filtered[lower_mask])
        statistics.update(
            {
                "interval_mean": interval_mean,
                "interval_max": interval_max,
                "interval_min": interval_min,
                "threshold1": threshold1,
                "threshold2": threshold2,
                "upper_mean": upper_mean,
                "upper_var": upper_var,
                "lower_mean": lower_mean,
                "lower_var": lower_var,
                "upper_n": int(upper_mask.sum()),
                "lower_n": int(lower_mask.sum()),
            }
        )
        if not upper_mask.any():
            qc_flags.append("upper_no_points")
        elif upper_mask.sum() == 1:
            qc_flags.append("upper_variance_insufficient")
        if not lower_mask.any():
            qc_flags.append("lower_no_points")
        elif lower_mask.sum() == 1:
            qc_flags.append("lower_variance_insufficient")
    else:
        qc_flags.append("motion_interval_out_of_range")

    return {
        "time": time,
        "raw": raw,
        "filtered": filtered,
        "motion_start_s": float(motion_start_s),
        "pre_interval_s": (pre_start, pre_end),
        "motion_interval_s": (interval_start, interval_end),
        "masks": {
            "pre30": pre_mask,
            "interval": interval_mask,
            "upper": upper_mask,
            "lower": lower_mask,
        },
        "statistics": statistics,
        "qc_flags": qc_flags,
        "debug": {"median_filter_samples": filter_samples},
    }


def analyze_interval_sample(
    sensor_csv: str | Path,
    *,
    data_root: str | Path,
    fs: int = 100,
    config: IntervalAnalysisConfig | None = None,
) -> dict[str, Any]:
    """读取一个样本，检测运动起点并独立统计 Udeco2/Udeco3。"""

    config = config or IntervalAnalysisConfig()
    loaded = load_udeco_sample(sensor_csv, data_root=data_root, fs=fs, epsilon=config.epsilon)
    motion = detect_motion_stage(
        loaded.accx,
        loaded.accy,
        loaded.accz,
        fs=loaded.fs,
        config=_motion_analysis_config(config),
    )
    if motion["status"] == "ok" and np.isfinite(motion["motion_start_s"]):
        motion_start_s = float(motion["motion_start_s"])
        motion_start_source = "detected"
        sample_qc: list[str] = []
    else:
        motion_start_s = float(config.fallback_motion_start_sec)
        motion_start_source = f"fallback_{_format_seconds_token(motion_start_s)}s"
        sample_qc = [f"motion_detection_failed_{motion_start_source}"]

    if int(loaded.qc["InterpFlag"].sum()) > 0:
        sample_qc.append("sequence_interpolated")
    if int(loaded.qc["ud_denominator_near_zero"].sum()) > 0:
        sample_qc.append("ud_denominator_near_zero")

    channels: dict[str, dict[str, Any]] = {}
    for channel, values in (("Udeco2", loaded.udeco2), ("Udeco3", loaded.udeco3)):
        channel_result = compute_channel_statistics(
            values,
            loaded.time_s,
            motion_start_s,
            loaded.fs,
            config,
        )
        channels[channel] = channel_result
        sample_qc.extend(f"{channel}:{flag}" for flag in channel_result["qc_flags"])

    qc_text = ";".join(dict.fromkeys(sample_qc)) if sample_qc else "OK"
    return {
        "sample_name": loaded.sample_name,
        "subject": loaded.subject,
        "fs": loaded.fs,
        "channels": channels,
        "motion": motion,
        "motion_start_s": motion_start_s,
        "motion_start_source": motion_start_source,
        "motion_detection_status": motion["status"],
        "motion_detection_reason": motion.get("reason", ""),
        "QC_flag": qc_text,
        "input_qc": loaded.qc,
        "metadata": loaded.metadata,
    }


def analyze_interval_subject(
    data_root: str | Path,
    subject: str,
    *,
    fs: int = 100,
    config: IntervalAnalysisConfig | None = None,
) -> dict[str, Any]:
    """分析一个受试者的全部严格配对样本，单样本失败不终止批次。"""

    config = config or IntervalAnalysisConfig()
    pairs, discovery_issues = discover_subject_samples(data_root, subject)
    rows: list[dict[str, Any]] = []
    qc_rows = [
        {
            "sample_name": "",
            "file_name": issue.file_name,
            "category": issue.category,
            "reason": issue.reason,
        }
        for issue in discovery_issues
    ]
    sample_details: dict[str, dict[str, Any]] = {}
    for pair in pairs:
        try:
            result = analyze_interval_sample(pair.sensor_csv, data_root=data_root, fs=fs, config=config)
            sample_details[pair.sample_name] = result
            rows.append(_sample_result_row(result, pair.motion_type, pair.motion_index))
        except Exception as exc:
            message = f"{type(exc).__name__}: {exc}"
            rows.append(_error_result_row(pair.sample_name, subject, pair.motion_type, pair.motion_index, message))
            qc_rows.append(
                {
                    "sample_name": pair.sample_name,
                    "file_name": pair.sensor_csv.name,
                    "category": "analysis_error",
                    "reason": message,
                }
            )
    return {
        "subject": subject,
        "sample_results": pd.DataFrame(rows, columns=_sample_result_columns()),
        "qc": pd.DataFrame(qc_rows, columns=["sample_name", "file_name", "category", "reason"]),
        "sample_details": sample_details,
    }


def analyze_interval_batch(
    data_root: str | Path,
    subjects: Sequence[str],
    *,
    fs: int = 100,
    config: IntervalAnalysisConfig | None = None,
) -> dict[str, dict[str, Any]]:
    """分析多个受试者；缺失受试者目录也保留独立 sheet 所需结果。"""

    config = config or IntervalAnalysisConfig()
    batch: dict[str, dict[str, Any]] = {}
    for subject in subjects:
        try:
            batch[subject] = analyze_interval_subject(data_root, subject, fs=fs, config=config)
        except Exception as exc:
            batch[subject] = {
                "subject": subject,
                "sample_results": pd.DataFrame(columns=_sample_result_columns()),
                "qc": pd.DataFrame(
                    [
                        {
                            "sample_name": "",
                            "file_name": "",
                            "category": "subject_error",
                            "reason": f"{type(exc).__name__}: {exc}",
                        }
                    ]
                ),
                "sample_details": {},
            }
    return batch


def plot_interval_validation(sample_result: dict[str, Any]) -> Any:
    """绘制 Udeco2/Udeco3 两行单文件验证图。"""

    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch

    figure, axes = plt.subplots(2, 1, figsize=(15, 8), sharex=True, constrained_layout=True)
    for axis, channel in zip(axes, ("Udeco2", "Udeco3")):
        result = sample_result["channels"][channel]
        time = result["time"]
        pre_start, pre_end = result["pre_interval_s"]
        interval_start, interval_end = result["motion_interval_s"]
        upper = result["masks"]["upper"]
        lower = result["masks"]["lower"]
        stats = result["statistics"]

        axis.axvspan(pre_start, pre_end, color="#2A9D8F", alpha=0.12, zorder=0)
        axis.axvspan(interval_start, interval_end, color="#E9C46A", alpha=0.16, zorder=0)
        axis.plot(time, result["raw"], color="#9A9A9A", linewidth=0.55, alpha=0.75, zorder=1)
        axis.plot(time, result["filtered"], color="black", linewidth=0.8, zorder=2)
        axis.scatter(time[upper], result["filtered"][upper], s=9, color="#D62828", zorder=4)
        axis.scatter(time[lower], result["filtered"][lower], s=9, color="#277DA1", zorder=4)
        if np.isfinite(stats["threshold1"]):
            axis.hlines(stats["threshold1"], interval_start, interval_end, color="#D62828", linestyle="--")
        if np.isfinite(stats["threshold2"]):
            axis.hlines(stats["threshold2"], interval_start, interval_end, color="#277DA1", linestyle="--")
        axis.axvline(sample_result["motion_start_s"], color="#F77F00", linestyle="--", linewidth=1.1)
        axis.axvline(interval_start, color="#8D6E00", linestyle=":", linewidth=0.9)
        axis.axvline(interval_end, color="#8D6E00", linestyle=":", linewidth=0.9)
        axis.set_ylabel(channel)
        axis.grid(alpha=0.15)

    handles = [
        Line2D([0], [0], color="#9A9A9A", label="Raw Udeco"),
        Line2D([0], [0], color="black", label="Median filtered"),
        Patch(facecolor="#2A9D8F", alpha=0.18, label="A: first 30 s"),
        Patch(facecolor="#E9C46A", alpha=0.22, label="Motion +15–35 s"),
        Line2D([0], [0], marker="o", linestyle="", color="#D62828", label="B: above threshold 1"),
        Line2D([0], [0], marker="o", linestyle="", color="#277DA1", label="C: below threshold 2"),
    ]
    axes[0].legend(handles=handles, loc="upper right", ncol=3)
    axes[-1].set_xlabel("Time / s")
    figure.suptitle(
        f"{sample_result['sample_name']} | motion start={sample_result['motion_start_s']:.3f} s "
        f"({sample_result['motion_start_source']})"
    )
    return figure


def export_interval_excel(
    batch_results: dict[str, dict[str, Any]],
    output_path: str | Path,
) -> Path:
    """将每个受试者的样本宽表和发现/分析 QC 写入同一个 Excel。"""

    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for subject, result in batch_results.items():
            sheet_name = _safe_sheet_name(subject)
            sample_results = result.get("sample_results", pd.DataFrame(columns=_sample_result_columns()))
            qc = result.get("qc", pd.DataFrame(columns=["sample_name", "file_name", "category", "reason"]))
            sample_results.to_excel(writer, sheet_name=sheet_name, startrow=2, index=False)
            qc_title_row = len(sample_results) + 6
            qc.to_excel(writer, sheet_name=sheet_name, startrow=qc_title_row, index=False)
            worksheet = writer.book[sheet_name]
            worksheet.cell(1, 1, "Part A: Sample-level interval statistics")
            worksheet.cell(qc_title_row, 1, "Part B: Skipped files / QC")
            title_fill = PatternFill("solid", fgColor="D9EAF7")
            for row in (1, qc_title_row):
                worksheet.cell(row, 1).font = Font(bold=True)
                worksheet.cell(row, 1).fill = title_fill
            worksheet.freeze_panes = "A4"
            for column_cells in worksheet.columns:
                letter = get_column_letter(column_cells[0].column)
                max_length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=8)
                worksheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 32)
    return output


def export_interval_means_excel(
    batch_results: dict[str, dict[str, Any]],
    output_path: str | Path,
) -> Path:
    """输出精简工作簿：每个样本仅保留中文运动类型和六个区域均值。"""

    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for subject, result in batch_results.items():
            source = result.get("sample_results", pd.DataFrame())
            compact = source.reindex(columns=COMPACT_MEAN_COLUMNS)
            compact["motion_type"] = compact["motion_type"].map(MOTION_TYPE_LABELS).fillna(compact["motion_type"])
            compact = compact.rename(columns={"motion_type": "运动类型"})
            sheet_name = _safe_sheet_name(subject)
            compact.to_excel(writer, sheet_name=sheet_name, index=False)
            worksheet = writer.book[sheet_name]
            header_fill = PatternFill("solid", fgColor="D9EAF7")
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            worksheet.freeze_panes = "B2"
            for column_cells in worksheet.columns:
                letter = get_column_letter(column_cells[0].column)
                max_length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=8)
                worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 28)
    return output


def _seconds_to_odd_samples(seconds: float, fs: int) -> int:
    samples = max(1, int(round(float(seconds) * fs)))
    return samples if samples % 2 == 1 else samples + 1


def _covers_half_open_interval(time: np.ndarray, start: float, end: float, fs: int, epsilon: float) -> bool:
    tolerance = max(float(epsilon), 0.51 / fs)
    return bool(time[0] <= start + tolerance and time[-1] + 1.0 / fs >= end - tolerance)


def _mean_and_sample_variance(values: np.ndarray) -> tuple[float, float]:
    finite = np.asarray(values, dtype=float)
    finite = finite[np.isfinite(finite)]
    if finite.size == 0:
        return float("nan"), float("nan")
    mean = float(np.mean(finite))
    variance = float(np.var(finite, ddof=1)) if finite.size >= 2 else float("nan")
    return mean, variance


def _motion_analysis_config(config: IntervalAnalysisConfig) -> AnalysisConfig:
    return AnalysisConfig(
        epsilon=config.epsilon,
        motion_calibration_sec=config.motion_calibration_sec,
        motion_window_sec=config.motion_window_sec,
        motion_step_sec=config.motion_step_sec,
        motion_threshold_sigma=config.motion_threshold_sigma,
        motion_transition_windows=config.motion_transition_windows,
    )


def _format_seconds_token(seconds: float) -> str:
    return str(int(seconds)) if float(seconds).is_integer() else f"{seconds:g}"


def _sample_result_columns() -> list[str]:
    identity = [
        "sample_name",
        "subject",
        "motion_type",
        "motion_index",
        "analysis_status",
        "error_message",
        "motion_detection_status",
        "motion_detection_reason",
        "motion_start_s",
        "motion_start_source",
        "motion_interval_start_s",
        "motion_interval_end_s",
        "QC_flag",
    ]
    channel_columns = [
        f"{channel}_{name}"
        for channel in ("Udeco2", "Udeco3")
        for name in (*CORE_STATISTICS, *DIAGNOSTIC_STATISTICS)
    ]
    return [*identity, *channel_columns]


def _sample_result_row(result: dict[str, Any], motion_type: str, motion_index: int) -> dict[str, Any]:
    first_channel = result["channels"]["Udeco2"]
    interval_start, interval_end = first_channel["motion_interval_s"]
    row: dict[str, Any] = {
        "sample_name": result["sample_name"],
        "subject": result["subject"],
        "motion_type": motion_type,
        "motion_index": motion_index,
        "analysis_status": "ok",
        "error_message": "",
        "motion_detection_status": result["motion_detection_status"],
        "motion_detection_reason": result["motion_detection_reason"],
        "motion_start_s": result["motion_start_s"],
        "motion_start_source": result["motion_start_source"],
        "motion_interval_start_s": interval_start,
        "motion_interval_end_s": interval_end,
        "QC_flag": result["QC_flag"],
    }
    for channel, channel_result in result["channels"].items():
        for name, value in channel_result["statistics"].items():
            row[f"{channel}_{name}"] = value
    return row


def _error_result_row(
    sample_name: str,
    subject: str,
    motion_type: str,
    motion_index: int,
    message: str,
) -> dict[str, Any]:
    row = {column: float("nan") for column in _sample_result_columns()}
    row.update(
        {
            "sample_name": sample_name,
            "subject": subject,
            "motion_type": motion_type,
            "motion_index": motion_index,
            "analysis_status": "error",
            "error_message": message,
            "motion_detection_status": "error",
            "motion_detection_reason": "",
            "motion_start_source": "",
            "QC_flag": "analysis_error",
        }
    )
    return row


def _safe_sheet_name(name: str) -> str:
    invalid = set('[]:*?/\\')
    cleaned = "".join("_" if char in invalid else char for char in str(name))[:31]
    return cleaned or "subject"


__all__ = [
    "IntervalAnalysisConfig",
    "compute_channel_statistics",
    "analyze_interval_sample",
    "analyze_interval_subject",
    "analyze_interval_batch",
    "plot_interval_validation",
    "export_interval_excel",
    "export_interval_means_excel",
]
