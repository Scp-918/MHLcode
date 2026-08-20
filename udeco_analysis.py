"""Udeco2/Udeco3 自适应状态分析工具。

本模块独立复现 change6 仓库中与 Pydisplay CSV、共享 ``ck.mat`` 标定和
100 Hz 时间轴相关的最小读取契约，不依赖或修改参考仓库。
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from scipy.io import loadmat
from scipy.ndimage import median_filter
from scipy.signal import butter, filtfilt, find_peaks


FS_HZ = 100
LEGAL_MOTION_TYPES = ("write", "gripper", "run", "rope")
PYDISPLAY_SENSOR_COLUMNS = (
    "frame_seq",
    "absolute_seq_u64",
    "segment_id",
    "sample_seq",
    "PPG_G",
    "PPG_R",
    "PPG_IR",
    "ACC_X",
    "ACC_Y",
    "ACC_Z",
    "GYRO_X",
    "GYRO_Y",
    "GYRO_Z",
    "Uh1",
    "Uh2",
    "Uh3",
    "Uh4",
    "Uc1",
    "Uc2",
    "Uc3",
    "Uc4",
    "UD1",
    "UD2",
    "parser_valid",
)

_FILE_RE = re.compile(
    r"^(?P<subject>.+)_(?P<motion>write|gripper|run|rope)_"
    r"(?P<index>\d+)_(?P<kind>sensor|HRdata)\.csv$"
)


@dataclass(frozen=True)
class Calibration:
    """Udeco 标定公式使用的通道 2/3 系数。"""

    source_path: Path
    c2: float
    c3: float


@dataclass(frozen=True)
class SamplePair:
    """严格配对的一份传感器和参考心率文件。"""

    sample_name: str
    subject: str
    motion_type: str
    motion_index: int
    sensor_csv: Path
    hr_csv: Path


@dataclass(frozen=True)
class DiscoveryIssue:
    """不能进入批处理的文件及其原因。"""

    file_name: str
    file_path: Path
    category: str
    reason: str


@dataclass
class LoadedSample:
    """经过序号修复、但未进行状态识别滤波的单样本数据。"""

    sample_name: str
    subject: str
    fs: int
    time_s: np.ndarray
    udeco2: np.ndarray
    udeco3: np.ndarray
    accx: np.ndarray
    accy: np.ndarray
    accz: np.ndarray
    qc: pd.DataFrame
    metadata: dict[str, Any]


@dataclass(frozen=True)
class AnalysisConfig:
    """状态识别参数；所有时间参数均使用秒。"""

    median_filter_sec: float = 0.03
    initial_m_sec: float = 30.0
    change_window_sec: float = 0.03
    change_threshold_sigma: float = 5.0
    change_peak_merge_sec: float = 0.04
    min_platform_sec: float = 0.06
    state_separation_sigma: float = 4.0
    m_anchor_sigma: float = 2.5
    m_baseline_smooth_sec: float = 10.0
    m_baseline_max_iter: int = 5
    enter_sigma: float = 4.0
    exit_sigma: float = 2.0
    use_gmm: bool = False
    epsilon: float = 1e-9
    motion_calibration_sec: float = 30.0
    motion_window_sec: float = 8.0
    motion_step_sec: float = 1.0
    motion_threshold_sigma: float = 3.0
    motion_transition_windows: int = 6

    def __post_init__(self) -> None:
        positive = {
            "median_filter_sec": self.median_filter_sec,
            "initial_m_sec": self.initial_m_sec,
            "change_window_sec": self.change_window_sec,
            "change_threshold_sigma": self.change_threshold_sigma,
            "change_peak_merge_sec": self.change_peak_merge_sec,
            "min_platform_sec": self.min_platform_sec,
            "state_separation_sigma": self.state_separation_sigma,
            "m_anchor_sigma": self.m_anchor_sigma,
            "m_baseline_smooth_sec": self.m_baseline_smooth_sec,
            "enter_sigma": self.enter_sigma,
            "exit_sigma": self.exit_sigma,
            "epsilon": self.epsilon,
        }
        invalid = [name for name, value in positive.items() if not np.isfinite(value) or value <= 0]
        if invalid:
            raise ValueError(f"Config values must be finite and positive: {', '.join(invalid)}")
        if self.enter_sigma <= self.exit_sigma:
            raise ValueError("enter_sigma must be greater than exit_sigma")
        if int(self.m_baseline_max_iter) < 1:
            raise ValueError("m_baseline_max_iter must be at least 1")


def load_calibration(data_root: str | Path) -> Calibration:
    """读取并验证 ``DATA_ROOT/ck.mat`` 中的 c2、c3。"""

    path = Path(data_root).resolve() / "ck.mat"
    if not path.is_file():
        raise FileNotFoundError(f"Calibration file not found: {path}")
    try:
        payload = loadmat(path, simplify_cells=True)
    except Exception as exc:  # pragma: no cover - scipy 提供具体底层错误
        raise ValueError(f"Cannot parse calibration MAT file {path}: {exc}") from exc
    c_result = _cell_sequence(payload, "cResult", minimum=3)
    c2_entry = _mat_struct(c_result[1], "cResult{2}")
    c3_entry = _mat_struct(c_result[2], "cResult{3}")
    _expect_integer(c2_entry, "channel", 2, "cResult{2}.channel")
    _expect_integer(c3_entry, "channel", 3, "cResult{3}.channel")
    return Calibration(
        source_path=path,
        c2=_finite_scalar(c2_entry, "c", "cResult{2}.c"),
        c3=_finite_scalar(c3_entry, "c", "cResult{3}.c"),
    )


def discover_subject_samples(
    data_root: str | Path,
    subject: str,
) -> tuple[list[SamplePair], list[DiscoveryIssue]]:
    """按 change6 的严格文件名契约发现一个受试者的成对样本。"""

    subject_dir = Path(data_root) / subject
    if not subject_dir.is_dir():
        raise NotADirectoryError(f"Subject directory not found: {subject_dir}")
    accepted: dict[tuple[str, int, str], list[Path]] = {}
    issues: list[DiscoveryIssue] = []
    for path in sorted(subject_dir.glob("*.csv"), key=lambda item: item.name):
        match = _FILE_RE.fullmatch(path.name)
        if match is None:
            issues.append(_issue(path, "invalid_name", "filename does not match the subject data contract"))
            continue
        if match.group("subject") != subject:
            issues.append(_issue(path, "subject_mismatch", "filename subject differs from directory"))
            continue
        if not _has_data_row(path):
            issues.append(_issue(path, "empty_data", "CSV contains no data rows"))
            continue
        key = (match.group("motion"), int(match.group("index")), match.group("kind"))
        accepted.setdefault(key, []).append(path)

    duplicate_keys = {key for key, paths in accepted.items() if len(paths) > 1}
    for key in sorted(duplicate_keys):
        for path in accepted[key]:
            issues.append(_issue(path, "duplicate", "duplicate normalized motion/index/kind"))

    order = {motion: index for index, motion in enumerate(LEGAL_MOTION_TYPES)}
    sample_keys = sorted(
        {(motion, index) for motion, index, _ in accepted},
        key=lambda item: (order[item[0]], item[1]),
    )
    pairs: list[SamplePair] = []
    for motion, index in sample_keys:
        sensor_key = (motion, index, "sensor")
        hr_key = (motion, index, "HRdata")
        sensors = [] if sensor_key in duplicate_keys else accepted.get(sensor_key, [])
        hrs = [] if hr_key in duplicate_keys else accepted.get(hr_key, [])
        if len(sensors) == 1 and len(hrs) == 1:
            sample_name = f"{subject}_{motion}_{index}"
            pairs.append(SamplePair(sample_name, subject, motion, index, sensors[0], hrs[0]))
        else:
            if len(sensors) == 1:
                issues.append(_issue(sensors[0], "missing_pair", "missing matching HRdata CSV"))
            if len(hrs) == 1:
                issues.append(_issue(hrs[0], "missing_pair", "missing matching sensor CSV"))
    return pairs, sorted(issues, key=lambda item: item.file_name)


def load_udeco_sample(
    sensor_csv: str | Path,
    *,
    data_root: str | Path,
    fs: int = FS_HZ,
    epsilon: float = 1e-9,
) -> LoadedSample:
    """读取一份 Pydisplay CSV，修复序号并计算 Udeco2/Udeco3。"""

    path = Path(sensor_csv)
    if not path.is_file():
        raise FileNotFoundError(f"Sensor CSV not found: {path}")
    if int(fs) <= 0:
        raise ValueError("fs must be positive")
    calibration = load_calibration(data_root)
    raw = pd.read_csv(path)
    missing = [column for column in PYDISPLAY_SENSOR_COLUMNS if column not in raw.columns]
    if missing:
        raise KeyError(f"Missing required Pydisplay columns: {', '.join(missing)}")
    frame, qc, metadata = _repair_pydisplay_sequence(raw)

    numeric_columns = ("Uh2", "Uh3", "Uc2", "Uc3", "ACC_X", "ACC_Y", "ACC_Z")
    for column in numeric_columns:
        values = pd.to_numeric(frame[column], errors="coerce")
        if not np.isfinite(values.to_numpy(dtype=float)).any():
            raise ValueError(f"Required channel {column!r} has no finite values")
        frame[column] = values.interpolate(limit_direction="both")

    den2 = calibration.c2 - frame["Uc2"].to_numpy(dtype=float)
    den3 = calibration.c3 - frame["Uc3"].to_numpy(dtype=float)
    near2 = np.isfinite(den2) & (np.abs(den2) <= float(epsilon))
    near3 = np.isfinite(den3) & (np.abs(den3) <= float(epsilon))
    safe2 = den2.copy()
    safe3 = den3.copy()
    safe2[near2] = np.nan
    safe3[near3] = np.nan
    udeco2 = (frame["Uh2"].to_numpy(dtype=float) - frame["Uc2"].to_numpy(dtype=float)) / safe2
    udeco3 = (frame["Uh3"].to_numpy(dtype=float) - frame["Uc3"].to_numpy(dtype=float)) / safe3
    udeco2 = _interpolate_finite(udeco2, "Udeco2")
    udeco3 = _interpolate_finite(udeco3, "Udeco3")
    qc = qc.copy()
    qc["ud_denominator_near_zero"] = (near2 | near3).astype(int)

    sample_name = path.name.removesuffix("_sensor.csv")
    match = _FILE_RE.fullmatch(path.name)
    subject = match.group("subject") if match else path.parent.name
    metadata.update(
        {
            "sensor_csv": str(path.resolve()),
            "calibration_path": str(calibration.source_path),
            "calibration_c2": calibration.c2,
            "calibration_c3": calibration.c3,
            "fs_hz": int(fs),
            "udeco_mapping": "Udeco2=ud1(Uh2/Uc2), Udeco3=ud2(Uh3/Uc3)",
        }
    )
    n = len(frame)
    return LoadedSample(
        sample_name=sample_name,
        subject=subject,
        fs=int(fs),
        time_s=np.arange(n, dtype=float) / float(fs),
        udeco2=udeco2,
        udeco3=udeco3,
        accx=frame["ACC_X"].to_numpy(dtype=float),
        accy=frame["ACC_Y"].to_numpy(dtype=float),
        accz=frame["ACC_Z"].to_numpy(dtype=float),
        qc=qc,
        metadata=metadata,
    )


def robust_sigma(values: np.ndarray, *, epsilon: float = 1e-9) -> float:
    """用 ``1.4826 × MAD`` 估计噪声尺度，并提供数值下限。"""

    arr = np.asarray(values, dtype=float)
    arr = arr[np.isfinite(arr)]
    if arr.size == 0:
        return float(epsilon)
    center = float(np.median(arr))
    mad = float(np.median(np.abs(arr - center)))
    return max(float(epsilon), 1.4826 * mad)


def estimate_initial_m(
    signal: np.ndarray,
    *,
    fs: int,
    initial_m_sec: float,
    epsilon: float = 1e-9,
) -> tuple[float, float, int]:
    """由记录开头估计 ``(M0, sigma0, 使用样本数)``。"""

    arr = np.asarray(signal, dtype=float)
    if arr.size == 0 or not np.isfinite(arr).all():
        raise ValueError("signal must contain finite values")
    n_initial = min(arr.size, max(1, int(round(float(initial_m_sec) * int(fs)))))
    initial = arr[:n_initial]
    m0 = float(np.median(initial))
    return m0, robust_sigma(initial, epsilon=epsilon), n_initial


def detect_change_points(
    signal: np.ndarray,
    *,
    fs: int,
    sigma0: float,
    config: AnalysisConfig,
) -> tuple[np.ndarray, np.ndarray]:
    """用短时左右窗口中位数差检测并合并平台变化点。"""

    arr = np.asarray(signal, dtype=float)
    n = arr.size
    window = max(1, int(round(config.change_window_sec * int(fs))))
    score = np.zeros(n, dtype=float)
    scale = max(float(sigma0), config.epsilon)
    if n < 2 * window + 1:
        return np.asarray([], dtype=int), score
    for index in range(window, n - window):
        left = float(np.median(arr[index - window : index]))
        right = float(np.median(arr[index : index + window]))
        score[index] = abs(right - left) / scale
    peaks, _ = find_peaks(score, height=config.change_threshold_sigma)
    if peaks.size == 0:
        return np.asarray([], dtype=int), score

    # 将得分峰细化到局部最大相邻跃变处，使 segment 边界贴近真实阶跃。
    refined: list[int] = []
    local_delta = np.abs(np.diff(arr))
    for peak in peaks:
        lo = max(0, int(peak) - window)
        hi = min(local_delta.size, int(peak) + window)
        boundary = int(np.argmax(local_delta[lo:hi]) + lo + 1) if hi > lo else int(peak)
        refined.append(boundary)
    candidates = sorted(set(index for index in refined if 0 < index < n))
    merge_samples = max(1, int(round(config.change_peak_merge_sec * int(fs))))
    groups: list[list[int]] = []
    for index in candidates:
        if not groups or index - groups[-1][-1] > merge_samples:
            groups.append([index])
        else:
            groups[-1].append(index)
    merged = [max(group, key=lambda idx: score[min(max(idx, 0), n - 1)]) for group in groups]
    return np.asarray(merged, dtype=int), score


def build_segments(
    raw: np.ndarray,
    filtered: np.ndarray,
    time: np.ndarray,
    change_points: np.ndarray,
    *,
    fs: int,
) -> pd.DataFrame:
    """以变化点为自然边界，计算每个平台的稳健与常规统计量。"""

    raw_arr = np.asarray(raw, dtype=float)
    filt_arr = np.asarray(filtered, dtype=float)
    time_arr = np.asarray(time, dtype=float)
    if not (raw_arr.size == filt_arr.size == time_arr.size):
        raise ValueError("raw, filtered and time must have equal length")
    boundaries = [0, *sorted(set(int(x) for x in change_points if 0 < int(x) < raw_arr.size)), raw_arr.size]
    records: list[dict[str, Any]] = []
    for segment_id, (start, end) in enumerate(zip(boundaries[:-1], boundaries[1:])):
        values = raw_arr[start:end]
        filtered_values = filt_arr[start:end]
        median = float(np.median(values))
        records.append(
            {
                "segment_id": segment_id,
                "start_idx": start,
                "end_idx": end,
                "start_time": float(time_arr[start]),
                "end_time": float(time_arr[end - 1] + 1.0 / fs),
                "duration": float((end - start) / fs),
                "median": median,
                "mean": float(np.mean(values)),
                "std": float(np.std(values, ddof=1)) if values.size > 1 else 0.0,
                "MAD": float(np.median(np.abs(values - median))),
                "filtered_median": float(np.median(filtered_values)),
            }
        )
    return pd.DataFrame.from_records(records)


def classify_segments(
    segments: pd.DataFrame,
    *,
    raw: np.ndarray,
    m0: float,
    sigma0: float,
    initial_n: int,
    config: AnalysisConfig,
) -> pd.DataFrame:
    """用 M0 锚点给自然 segment 生成第一次 M/H/L 粗分类。"""

    result = segments.copy()
    threshold = config.state_separation_sigma * max(float(sigma0), config.epsilon)
    labels: list[str] = []
    raw_arr = np.asarray(raw, dtype=float)
    for row in result.itertuples(index=False):
        overlap_end = min(int(row.end_idx), int(initial_n))
        initial_center = (
            float(np.median(raw_arr[int(row.start_idx) : overlap_end]))
            if int(row.start_idx) < overlap_end
            else float("nan")
        )
        center = initial_center if np.isfinite(initial_center) else float(row.median)
        distance = center - float(m0)
        if abs(distance) <= threshold:
            labels.append("M")
        else:
            labels.append("H" if distance > 0 else "L")
    if config.use_gmm and len(result) >= 2:
        labels = _gmm_assisted_labels(result, labels, m0=m0, sigma0=sigma0, config=config)
    result["initial_label"] = labels
    return result


def estimate_dynamic_m(
    filtered: np.ndarray,
    segments: pd.DataFrame,
    *,
    fs: int,
    sigma0: float,
    config: AnalysisConfig,
) -> tuple[np.ndarray, np.ndarray, pd.DataFrame]:
    """迭代扩展可信 M segment，并估计缓慢变化的 M(t)。"""

    arr = np.asarray(filtered, dtype=float)
    work = segments.copy()
    trusted_ids = set(work.loc[work["initial_label"] == "M", "segment_id"].astype(int))
    if not trusted_ids:
        trusted_ids.add(int(work.iloc[0]["segment_id"]))
    m_t = np.full(arr.size, float(np.median(arr)), dtype=float)
    for _ in range(int(config.m_baseline_max_iter)):
        trusted_mask = _segment_sample_mask(arr.size, work, trusted_ids)
        m_t = _baseline_from_mask(arr, trusted_mask, fs=fs, smooth_sec=config.m_baseline_smooth_sec)
        trusted_residual = arr[trusted_mask] - m_t[trusted_mask]
        scale = max(float(sigma0), robust_sigma(trusted_residual, epsilon=config.epsilon))
        residual_medians = []
        for row in work.itertuples(index=False):
            start, end = int(row.start_idx), int(row.end_idx)
            residual_medians.append(float(np.median(arr[start:end] - m_t[start:end])))
        work["residual_median"] = residual_medians
        expanded = set(
            work.loc[
                work["residual_median"].abs() <= config.m_anchor_sigma * scale,
                "segment_id",
            ].astype(int)
        ) | trusted_ids
        if expanded == trusted_ids:
            break
        trusted_ids = expanded
    trusted_mask = _segment_sample_mask(arr.size, work, trusted_ids)
    m_t = _baseline_from_mask(arr, trusted_mask, fs=fs, smooth_sec=config.m_baseline_smooth_sec)
    work["trusted_m"] = work["segment_id"].astype(int).isin(trusted_ids)
    work["residual_median"] = [
        float(np.median(arr[int(row.start_idx) : int(row.end_idx)] - m_t[int(row.start_idx) : int(row.end_idx)]))
        for row in work.itertuples(index=False)
    ]
    return m_t, trusted_mask, work


def apply_hysteresis(
    residual_medians: np.ndarray,
    durations: np.ndarray,
    *,
    sigma_m: float,
    config: AnalysisConfig,
) -> tuple[np.ndarray, list[str]]:
    """在 segment 层执行进入/退出阈值不同的 H/M/L 滞回。"""

    residuals = np.asarray(residual_medians, dtype=float)
    durations_arr = np.asarray(durations, dtype=float)
    enter = config.enter_sigma * max(float(sigma_m), config.epsilon)
    exit_threshold = config.exit_sigma * max(float(sigma_m), config.epsilon)
    state = "M"
    labels: list[str] = []
    qc_flags: list[str] = []
    for residual, duration in zip(residuals, durations_arr):
        allow_new_state = duration + config.epsilon >= config.min_platform_sec
        previous = state
        if state == "M":
            if allow_new_state and residual > enter:
                state = "H"
            elif allow_new_state and residual < -enter:
                state = "L"
        elif state == "H":
            if allow_new_state and residual < -enter:
                state = "L"
            elif residual < exit_threshold:
                state = "M"
        else:  # L
            if allow_new_state and residual > enter:
                state = "H"
            elif residual > -exit_threshold:
                state = "M"
        if previous == "H" and state == "L":
            qc_flags.append("direct_H_to_L")
        elif previous == "L" and state == "H":
            qc_flags.append("direct_L_to_H")
        labels.append(state)
    return np.asarray(labels, dtype=object), list(dict.fromkeys(qc_flags))


def classify_udeco(
    raw: np.ndarray,
    *,
    time: np.ndarray | None = None,
    fs: int = FS_HZ,
    config: AnalysisConfig | None = None,
) -> dict[str, Any]:
    """完整执行单路 Udeco 的变化点、动态 M 与滞回分类。"""

    config = config or AnalysisConfig()
    raw_arr = np.asarray(raw, dtype=float).ravel()
    if raw_arr.size < 3 or not np.isfinite(raw_arr).all():
        raise ValueError("raw must contain at least three finite samples")
    if int(fs) <= 0:
        raise ValueError("fs must be positive")
    time_arr = np.arange(raw_arr.size, dtype=float) / float(fs) if time is None else np.asarray(time, dtype=float)
    if time_arr.shape != raw_arr.shape or not np.isfinite(time_arr).all():
        raise ValueError("time must be finite and have the same shape as raw")
    filter_samples = _odd_samples(config.median_filter_sec, fs)
    filtered = median_filter(raw_arr, size=filter_samples, mode="nearest")
    m0, sigma0, initial_n = estimate_initial_m(
        filtered,
        fs=fs,
        initial_m_sec=config.initial_m_sec,
        epsilon=config.epsilon,
    )
    change_points, score = detect_change_points(filtered, fs=fs, sigma0=sigma0, config=config)
    segments = build_segments(raw_arr, filtered, time_arr, change_points, fs=fs)
    segments = classify_segments(
        segments,
        raw=raw_arr,
        m0=m0,
        sigma0=sigma0,
        initial_n=initial_n,
        config=config,
    )
    m_t, trusted_mask, segments = estimate_dynamic_m(
        filtered,
        segments,
        fs=fs,
        sigma0=sigma0,
        config=config,
    )
    residual = filtered - m_t
    sigma_m = max(sigma0, robust_sigma(residual[trusted_mask], epsilon=config.epsilon))
    segment_labels, qc_flags = apply_hysteresis(
        segments["residual_median"].to_numpy(dtype=float),
        segments["duration"].to_numpy(dtype=float),
        sigma_m=sigma_m,
        config=config,
    )
    segments["final_label"] = segment_labels
    segments["platform_ok"] = segments["duration"] + config.epsilon >= config.min_platform_sec
    labels = np.full(raw_arr.size, "M", dtype=object)
    for row in segments.itertuples(index=False):
        labels[int(row.start_idx) : int(row.end_idx)] = str(row.final_label)
    if initial_n < int(round(config.initial_m_sec * fs)):
        qc_flags.append("initial_window_short")
    statistics = _state_statistics(raw_arr, labels, filtered - m_t, fs=fs, m0=m0, sigma0=sigma0, sigma_m=sigma_m)
    statistics["QC_flag"] = ";".join(dict.fromkeys(qc_flags)) if qc_flags else "OK"
    return {
        "time": time_arr,
        "raw": raw_arr,
        "filtered": filtered,
        "labels": labels,
        "M_t": m_t,
        "residual": residual,
        "change_points": change_points,
        "segments": segments,
        "statistics": statistics,
        "debug": {
            "change_score": score,
            "change_threshold_sigma": config.change_threshold_sigma,
            "sigma_m": sigma_m,
            "enter_threshold": config.enter_sigma * sigma_m,
            "exit_threshold": config.exit_sigma * sigma_m,
            "filter_samples": filter_samples,
            "trusted_m_mask": trusted_mask,
        },
    }


def detect_motion_stage(
    accx: np.ndarray,
    accy: np.ndarray,
    accz: np.ndarray,
    *,
    fs: int = FS_HZ,
    config: AnalysisConfig | None = None,
) -> dict[str, Any]:
    """沿用参考仓库的 ACC 合模长 STD 规则识别运动起止。"""

    config = config or AnalysisConfig()
    axes = [_safe_bandpass(np.asarray(axis, dtype=float), fs, 0.5, 10.0) for axis in (accx, accy, accz)]
    if len({axis.size for axis in axes}) != 1 or axes[0].size == 0:
        return _motion_error("ACC channels must have equal non-zero length")
    magnitude = np.sqrt(sum(axis**2 for axis in axes))
    calibration_n = min(magnitude.size, int(round(config.motion_calibration_sec * fs)))
    if calibration_n < 2:
        return _motion_error("not enough samples for motion calibration")
    baseline_std = float(np.std(magnitude[:calibration_n], ddof=1))
    threshold = config.motion_threshold_sigma * max(baseline_std, config.epsilon)
    window_n = max(2, int(round(config.motion_window_sec * fs)))
    step_n = max(1, int(round(config.motion_step_sec * fs)))
    if magnitude.size < window_n:
        return _motion_error("signal shorter than one motion window", threshold=threshold)
    starts = np.arange(0, magnitude.size - window_n + 1, step_n, dtype=int)
    window_std = np.asarray(
        [np.std(magnitude[start : start + window_n], ddof=1) for start in starts],
        dtype=float,
    )
    flags = window_std > threshold
    run = int(config.motion_transition_windows)
    start_index = _find_boolean_transition(flags, before=False, after=True, run=run, start_at=0)
    if start_index is None:
        return _motion_error(
            f"cannot find {run} rest windows followed by {run} motion windows",
            threshold=threshold,
            starts=starts / fs,
            window_std=window_std,
            flags=flags,
        )
    end_index = _find_boolean_transition(flags, before=True, after=False, run=run, start_at=start_index)
    if end_index is None:
        return _motion_error(
            f"cannot find {run} motion windows followed by {run} rest windows",
            threshold=threshold,
            starts=starts / fs,
            window_std=window_std,
            flags=flags,
        )
    return {
        "status": "ok",
        "reason": "",
        "motion_start_s": float(starts[start_index] / fs),
        "motion_end_s": float(starts[end_index] / fs),
        "motion_threshold": float(threshold),
        "window_starts_s": starts.astype(float) / fs,
        "window_centers_s": starts.astype(float) / fs + config.motion_window_sec / 2.0,
        "window_std": window_std,
        "motion_flags": flags,
    }


def analyze_sample(
    sensor_csv: str | Path,
    *,
    data_root: str | Path,
    fs: int = FS_HZ,
    config: AnalysisConfig | None = None,
) -> dict[str, Any]:
    """读取一个文件，并让 Udeco2/Udeco3 独立执行同一分类函数。"""

    config = config or AnalysisConfig()
    loaded = load_udeco_sample(sensor_csv, data_root=data_root, fs=fs, epsilon=config.epsilon)
    motion = detect_motion_stage(loaded.accx, loaded.accy, loaded.accz, fs=loaded.fs, config=config)
    channels: dict[str, dict[str, Any]] = {}
    for name, values in (("Udeco2", loaded.udeco2), ("Udeco3", loaded.udeco3)):
        channel_result = classify_udeco(values, time=loaded.time_s, fs=loaded.fs, config=config)
        statistics = channel_result["statistics"]
        statistics.update(
            {
                "sample_name": loaded.sample_name,
                "subject": loaded.subject,
                "channel": name,
                "analysis_status": "ok",
                "error_message": "",
                "motion_start_s": motion["motion_start_s"],
                "motion_end_s": motion["motion_end_s"],
            }
        )
        extra_qc: list[str] = []
        if int(loaded.qc["ud_denominator_near_zero"].sum()) > 0:
            extra_qc.append("ud_denominator_near_zero")
        if int(loaded.qc["InterpFlag"].sum()) > 0:
            extra_qc.append("sequence_interpolated")
        if motion["status"] != "ok":
            extra_qc.append("motion_detection_failed")
        current_qc = [] if statistics["QC_flag"] == "OK" else statistics["QC_flag"].split(";")
        combined = list(dict.fromkeys([*current_qc, *extra_qc]))
        statistics["QC_flag"] = ";".join(combined) if combined else "OK"
        channels[name] = channel_result
    return {
        "sample_name": loaded.sample_name,
        "subject": loaded.subject,
        "fs": loaded.fs,
        "time": loaded.time_s,
        "channels": channels,
        "motion": motion,
        "input_qc": loaded.qc,
        "metadata": loaded.metadata,
    }


def analyze_subject(
    data_root: str | Path,
    subject: str,
    *,
    fs: int = FS_HZ,
    config: AnalysisConfig | None = None,
) -> dict[str, Any]:
    """分析一个受试者的全部严格配对样本；单样本错误不终止批次。"""

    config = config or AnalysisConfig()
    pairs, discovery_issues = discover_subject_samples(data_root, subject)
    result_rows: list[dict[str, Any]] = []
    qc_rows = [
        {
            "sample_name": "",
            "file_name": issue.file_name,
            "category": issue.category,
            "reason": issue.reason,
        }
        for issue in discovery_issues
    ]
    sample_details: dict[str, dict[str, Any]] = {}
    for pair in pairs:
        try:
            sample_result = analyze_sample(pair.sensor_csv, data_root=data_root, fs=fs, config=config)
            sample_details[pair.sample_name] = sample_result
            for channel_name, channel_result in sample_result["channels"].items():
                row = dict(channel_result["statistics"])
                row.update({"motion_type": pair.motion_type, "motion_index": pair.motion_index})
                result_rows.append(row)
        except Exception as exc:
            message = f"{type(exc).__name__}: {exc}"
            qc_rows.append(
                {
                    "sample_name": pair.sample_name,
                    "file_name": pair.sensor_csv.name,
                    "category": "analysis_error",
                    "reason": message,
                }
            )
            for channel_name in ("Udeco2", "Udeco3"):
                result_rows.append(
                    _error_statistics_row(
                        sample_name=pair.sample_name,
                        subject=subject,
                        channel=channel_name,
                        motion_type=pair.motion_type,
                        motion_index=pair.motion_index,
                        message=message,
                    )
                )
    sample_results = pd.DataFrame(result_rows, columns=_sample_result_columns())
    qc = pd.DataFrame(qc_rows, columns=["sample_name", "file_name", "category", "reason"])
    return {
        "subject": subject,
        "sample_results": sample_results,
        "summary": summarize_subject(sample_results),
        "qc": qc,
        "sample_details": sample_details,
    }


def analyze_all_subjects(
    data_root: str | Path,
    subjects: list[str] | tuple[str, ...],
    *,
    fs: int = FS_HZ,
    config: AnalysisConfig | None = None,
) -> dict[str, dict[str, Any]]:
    """依次运行全部受试者，并为缺失目录保留受试者级 QC。"""

    config = config or AnalysisConfig()
    batch: dict[str, dict[str, Any]] = {}
    for subject in subjects:
        try:
            batch[subject] = analyze_subject(data_root, subject, fs=fs, config=config)
        except Exception as exc:
            empty = pd.DataFrame(columns=_sample_result_columns())
            batch[subject] = {
                "subject": subject,
                "sample_results": empty,
                "summary": summarize_subject(empty),
                "qc": pd.DataFrame(
                    [
                        {
                            "sample_name": "",
                            "file_name": "",
                            "category": "subject_error",
                            "reason": f"{type(exc).__name__}: {exc}",
                        }
                    ]
                ),
                "sample_details": {},
            }
    return batch


def summarize_subject(sample_results: pd.DataFrame) -> pd.DataFrame:
    """汇总各通道 H/M/L 样本均值的跨样本 mean、std 和有效 N。"""

    columns = ["channel", "state", "mean", "std", "N", "mean_std"]
    records: list[dict[str, Any]] = []
    frame = sample_results.copy()
    if "analysis_status" in frame.columns:
        frame = frame.loc[frame["analysis_status"] == "ok"]
    for channel in ("Udeco2", "Udeco3"):
        channel_frame = frame.loc[frame.get("channel", pd.Series(dtype=object)) == channel]
        for state in ("H", "M", "L"):
            source = channel_frame.get(f"{state}_mean", pd.Series(dtype=float))
            values = pd.to_numeric(source, errors="coerce").dropna()
            count = int(values.size)
            mean = float(values.mean()) if count else float("nan")
            std = float(values.std(ddof=1)) if count >= 2 else float("nan")
            formatted = f"{mean:.6g} ± {std:.6g}" if count >= 2 else f"{mean:.6g}" if count == 1 else "NaN"
            records.append(
                {"channel": channel, "state": state, "mean": mean, "std": std, "N": count, "mean_std": formatted}
            )
    return pd.DataFrame.from_records(records, columns=columns)


def export_excel(batch_results: dict[str, dict[str, Any]], output_path: str | Path) -> Path:
    """把每个受试者的样本、汇总和 QC 写入独立 Excel sheet。"""

    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for subject, result in batch_results.items():
            sheet_name = _safe_sheet_name(subject)
            wide = _wide_sample_results(result["sample_results"])
            summary = result.get("summary", summarize_subject(result["sample_results"]))
            qc = result.get("qc", pd.DataFrame(columns=["sample_name", "file_name", "category", "reason"]))
            part_a_row = 1
            table_a_row = 3
            wide.to_excel(writer, sheet_name=sheet_name, startrow=table_a_row - 1, index=False)
            part_b_row = table_a_row + len(wide) + 3
            summary.to_excel(writer, sheet_name=sheet_name, startrow=part_b_row + 1, index=False)
            part_c_row = part_b_row + len(summary) + 5
            qc.to_excel(writer, sheet_name=sheet_name, startrow=part_c_row + 1, index=False)
            worksheet = writer.book[sheet_name]
            worksheet.cell(part_a_row, 1, "Part A: Sample-level results")
            worksheet.cell(part_b_row + 1, 1, "Part B: Subject summary")
            worksheet.cell(part_c_row + 1, 1, "Part C: Skipped files / QC")
            title_fill = PatternFill("solid", fgColor="D9EAF7")
            for row in (part_a_row, part_b_row + 1, part_c_row + 1):
                worksheet.cell(row, 1).font = Font(bold=True)
                worksheet.cell(row, 1).fill = title_fill
            worksheet.freeze_panes = "A4"
            for column_cells in worksheet.columns:
                values = [str(cell.value) for cell in column_cells if cell.value is not None]
                width = min(45, max([len(value) for value in values], default=8) + 2)
                worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    return output


def plot_validation(
    sample_result: dict[str, Any],
    *,
    show_debug_info: bool = True,
) -> tuple[Any, Any | None]:
    """绘制固定两行主验证图，并可选返回 4×2 调试图。"""

    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch

    colors = {"H": "#E76F51", "M": "#2A9D8F", "L": "#457B9D"}
    channel_names = ("Udeco2", "Udeco3")
    main_fig, main_axes = plt.subplots(2, 1, figsize=(15, 8), sharex=True, constrained_layout=True)
    motion = sample_result["motion"]
    for axis, channel_name in zip(main_axes, channel_names):
        result = sample_result["channels"][channel_name]
        time = result["time"]
        for start, end, label in _contiguous_label_runs(time, result["labels"], sample_result["fs"]):
            axis.axvspan(start, end, color=colors[label], alpha=0.16, zorder=0)
        axis.plot(time, result["raw"], color="black", linewidth=0.8, label="Raw Udeco", zorder=4)
        axis.plot(time, result["M_t"], color="#264653", linewidth=1.2, linestyle="--", label="M(t)", zorder=3)
        enter = float(result["debug"]["enter_threshold"])
        axis.plot(time, result["M_t"] + enter, color="#E76F51", linewidth=0.7, linestyle=":", alpha=0.8)
        axis.plot(time, result["M_t"] - enter, color="#457B9D", linewidth=0.7, linestyle=":", alpha=0.8)
        for point in result["change_points"]:
            axis.axvline(time[int(point)], color="#6C757D", linewidth=0.5, alpha=0.5, zorder=1)
        if motion["status"] == "ok":
            axis.axvline(motion["motion_start_s"], color="#F4A261", linestyle="--", linewidth=1.1)
            axis.axvline(motion["motion_end_s"], color="#F4A261", linestyle="--", linewidth=1.1)
        axis.set_ylabel(channel_name)
        axis.grid(alpha=0.15)
        handles = [
            Line2D([0], [0], color="black", linewidth=1.0, label="Raw Udeco"),
            Line2D([0], [0], color="#264653", linestyle="--", label="M(t)"),
            *[Patch(facecolor=colors[state], alpha=0.22, label=state) for state in ("H", "M", "L")],
        ]
        axis.legend(handles=handles, loc="upper right", ncol=5)
    main_axes[-1].set_xlabel("Time / s")
    main_fig.suptitle(f"{sample_result['sample_name']} | Udeco state validation")

    if not show_debug_info:
        return main_fig, None
    debug_fig, debug_axes = plt.subplots(4, 2, figsize=(16, 11), sharex="col", constrained_layout=True)
    for column, channel_name in enumerate(channel_names):
        result = sample_result["channels"][channel_name]
        time = result["time"]
        debug_axes[0, column].plot(time, result["filtered"], color="#333333", linewidth=0.8, label="Filtered")
        debug_axes[0, column].plot(time, result["M_t"], color="#2A9D8F", linestyle="--", label="M(t)")
        for point in result["change_points"]:
            debug_axes[0, column].axvline(time[int(point)], color="#E9C46A", linewidth=0.6)
        debug_axes[0, column].legend(loc="upper right")
        debug_axes[0, column].set_title(channel_name)
        debug_axes[1, column].plot(time, result["debug"]["change_score"], color="#6D597A", linewidth=0.8)
        debug_axes[1, column].axhline(
            result["debug"].get("change_threshold_sigma", np.nan), color="#D62828", linestyle="--"
        )
        debug_axes[1, column].set_ylabel("S(t)")
        debug_axes[2, column].plot(time, result["residual"], color="#1D3557", linewidth=0.8)
        enter = float(result["debug"]["enter_threshold"])
        exit_threshold = float(result["debug"]["exit_threshold"])
        for value, style in ((enter, "--"), (-enter, "--"), (exit_threshold, ":"), (-exit_threshold, ":")):
            debug_axes[2, column].axhline(value, color="#D62828", linestyle=style, linewidth=0.7)
        debug_axes[2, column].set_ylabel("Residual")
        code = pd.Series(result["labels"]).map({"L": -1, "M": 0, "H": 1}).to_numpy(dtype=float)
        debug_axes[3, column].step(time, code, where="post", color="#264653")
        debug_axes[3, column].set_yticks([-1, 0, 1], labels=["L", "M", "H"])
        debug_axes[3, column].set_xlabel("Time / s")
        debug_axes[3, column].set_ylabel("Label")
        for row in range(4):
            debug_axes[row, column].grid(alpha=0.15)
    debug_fig.suptitle(f"{sample_result['sample_name']} | Debug diagnostics")
    return main_fig, debug_fig


def _safe_bandpass(signal: np.ndarray, fs: int, low_hz: float, high_hz: float) -> np.ndarray:
    """对 ACC 使用零相位带通；短序列安全退化为去均值。"""

    arr = np.asarray(signal, dtype=float)
    if arr.size == 0 or not np.isfinite(arr).all():
        return arr
    nyquist = float(fs) / 2.0
    high = min(float(high_hz), 0.45 * float(fs))
    low = max(float(low_hz), 1e-3)
    if arr.size < 8 or not (0 < low < high < nyquist):
        return arr - np.mean(arr)
    b, a = butter(4, [low / nyquist, high / nyquist], btype="bandpass")
    padlen = 3 * max(len(a), len(b))
    if arr.size <= padlen:
        return arr - np.mean(arr)
    try:
        return filtfilt(b, a, arr)
    except ValueError:
        return arr - np.mean(arr)


def _find_boolean_transition(
    flags: np.ndarray,
    *,
    before: bool,
    after: bool,
    run: int,
    start_at: int,
) -> int | None:
    for index in range(max(0, int(start_at)), len(flags) - 2 * int(run) + 1):
        if np.all(flags[index : index + run] == before) and np.all(
            flags[index + run : index + 2 * run] == after
        ):
            return index + run
    return None


def _motion_error(
    reason: str,
    *,
    threshold: float = float("nan"),
    starts: np.ndarray | None = None,
    window_std: np.ndarray | None = None,
    flags: np.ndarray | None = None,
) -> dict[str, Any]:
    starts_arr = np.asarray([] if starts is None else starts, dtype=float)
    return {
        "status": "error",
        "reason": reason,
        "motion_start_s": float("nan"),
        "motion_end_s": float("nan"),
        "motion_threshold": float(threshold),
        "window_starts_s": starts_arr,
        "window_centers_s": starts_arr,
        "window_std": np.asarray([] if window_std is None else window_std, dtype=float),
        "motion_flags": np.asarray([] if flags is None else flags, dtype=bool),
    }


def _contiguous_label_runs(time: np.ndarray, labels: np.ndarray, fs: int) -> list[tuple[float, float, str]]:
    """把样本标签压缩为绘图使用的连续时间区间。"""

    time_arr = np.asarray(time, dtype=float)
    labels_arr = np.asarray(labels, dtype=object)
    if labels_arr.size == 0:
        return []
    changes = np.flatnonzero(labels_arr[1:] != labels_arr[:-1]) + 1
    boundaries = [0, *changes.tolist(), labels_arr.size]
    return [
        (
            float(time_arr[start]),
            float(time_arr[end - 1] + 1.0 / fs),
            str(labels_arr[start]),
        )
        for start, end in zip(boundaries[:-1], boundaries[1:])
    ]


def _sample_result_columns() -> list[str]:
    return [
        "sample_name",
        "subject",
        "motion_type",
        "motion_index",
        "channel",
        "analysis_status",
        "error_message",
        "n_detected_states",
        "H_mean",
        "H_std",
        "M_mean",
        "M_std",
        "L_mean",
        "L_std",
        "H_total_duration",
        "M_total_duration",
        "L_total_duration",
        "transition_count",
        "M0",
        "sigma0",
        "sigma_M",
        "state_separation",
        "motion_start_s",
        "motion_end_s",
        "QC_flag",
    ]


def _error_statistics_row(
    *,
    sample_name: str,
    subject: str,
    channel: str,
    motion_type: str,
    motion_index: int,
    message: str,
) -> dict[str, Any]:
    row: dict[str, Any] = {column: float("nan") for column in _sample_result_columns()}
    row.update(
        {
            "sample_name": sample_name,
            "subject": subject,
            "motion_type": motion_type,
            "motion_index": int(motion_index),
            "channel": channel,
            "analysis_status": "error",
            "error_message": message,
            "QC_flag": "analysis_error",
        }
    )
    return row


def _wide_sample_results(sample_results: pd.DataFrame) -> pd.DataFrame:
    """将内部 sample×channel 长表转换为每个样本一行的 Excel 宽表。"""

    base_columns = ["sample_name", "subject", "motion_type", "motion_index"]
    if sample_results.empty:
        return pd.DataFrame(columns=base_columns)
    records: list[dict[str, Any]] = []
    for sample_name, group in sample_results.groupby("sample_name", sort=False, dropna=False):
        first = group.iloc[0]
        record = {column: first.get(column, "") for column in base_columns}
        record["sample_name"] = sample_name
        for _, row in group.iterrows():
            channel = str(row["channel"])
            for column in _sample_result_columns():
                if column not in {*base_columns, "channel"}:
                    record[f"{channel}_{column}"] = row.get(column, np.nan)
        records.append(record)
    return pd.DataFrame.from_records(records)


def _safe_sheet_name(subject: str) -> str:
    invalid = set('[]:*?/\\')
    cleaned = "".join("_" if char in invalid else char for char in str(subject)).strip()
    return (cleaned or "subject")[:31]


def _state_statistics(
    raw: np.ndarray,
    labels: np.ndarray,
    residual: np.ndarray,
    *,
    fs: int,
    m0: float,
    sigma0: float,
    sigma_m: float,
) -> dict[str, Any]:
    stats: dict[str, Any] = {}
    present: list[str] = []
    non_m_separation: list[float] = []
    for state in ("H", "M", "L"):
        mask = labels == state
        if mask.any():
            present.append(state)
            values = raw[mask]
            stats[f"{state}_mean"] = float(np.mean(values))
            stats[f"{state}_std"] = float(np.std(values, ddof=1)) if values.size > 1 else 0.0
            stats[f"{state}_total_duration"] = float(mask.sum() / fs)
            if state != "M":
                non_m_separation.append(abs(float(np.median(residual[mask]))) / max(sigma_m, np.finfo(float).eps))
        else:
            stats[f"{state}_mean"] = float("nan")
            stats[f"{state}_std"] = float("nan")
            stats[f"{state}_total_duration"] = float("nan")
    stats.update(
        {
            "n_detected_states": len(present),
            "transition_count": int(np.sum(labels[1:] != labels[:-1])),
            "M0": float(m0),
            "sigma0": float(sigma0),
            "sigma_M": float(sigma_m),
            "state_separation": min(non_m_separation) if non_m_separation else float("nan"),
        }
    )
    return stats


def _baseline_from_mask(signal: np.ndarray, mask: np.ndarray, *, fs: int, smooth_sec: float) -> np.ndarray:
    source = pd.Series(np.where(mask, signal, np.nan), dtype=float)
    filled = source.interpolate(limit_direction="both")
    if not filled.notna().any():
        filled[:] = float(np.median(signal))
    window = _odd_samples(smooth_sec, fs)
    return filled.rolling(window=window, center=True, min_periods=1).median().to_numpy(dtype=float)


def _segment_sample_mask(length: int, segments: pd.DataFrame, trusted_ids: set[int]) -> np.ndarray:
    mask = np.zeros(length, dtype=bool)
    for row in segments.itertuples(index=False):
        if int(row.segment_id) in trusted_ids:
            mask[int(row.start_idx) : int(row.end_idx)] = True
    return mask


def _odd_samples(seconds: float, fs: int) -> int:
    samples = max(1, int(round(float(seconds) * int(fs))))
    return samples if samples % 2 == 1 else samples + 1


def _gmm_assisted_labels(
    segments: pd.DataFrame,
    labels: list[str],
    *,
    m0: float,
    sigma0: float,
    config: AnalysisConfig,
) -> list[str]:
    """可选 GMM/BIC 辅助；默认关闭且不改变主流程依赖。"""

    from sklearn.mixture import GaussianMixture

    values = segments["median"].to_numpy(dtype=float).reshape(-1, 1)
    models = []
    for components in range(1, min(3, len(values)) + 1):
        model = GaussianMixture(n_components=components, random_state=0, reg_covar=config.epsilon)
        model.fit(values)
        models.append((model.bic(values), model))
    model = min(models, key=lambda item: item[0])[1]
    centers = model.means_.ravel()
    m_component = int(np.argmin(np.abs(centers - m0)))
    assignments = model.predict(values)
    threshold = config.state_separation_sigma * max(sigma0, config.epsilon)
    assisted: list[str] = []
    for existing, component in zip(labels, assignments):
        center = float(centers[component])
        if component == m_component or abs(center - m0) <= threshold:
            assisted.append("M")
        else:
            assisted.append("H" if center > m0 else "L")
    return assisted


def _repair_pydisplay_sequence(raw: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, int]]:
    """过滤无效行，并在同一 segment 内补齐 absolute sequence 缺口。"""

    input_rows = len(raw)
    empty_mask = raw.loc[:, list(PYDISPLAY_SENSOR_COLUMNS)].isna().all(axis=1)
    removed_empty = int(empty_mask.sum())
    work = raw.loc[~empty_mask].copy()
    valid = work["parser_valid"].map(_parser_valid_value).fillna(False).astype(bool)
    removed_invalid = int((~valid).sum())
    work = work.loc[valid].copy()
    if work.empty:
        raise ValueError("Sensor CSV has no usable rows after parser_valid filtering")
    for column in ("frame_seq", "absolute_seq_u64", "segment_id", "sample_seq"):
        values = pd.to_numeric(work[column], errors="coerce").to_numpy(dtype=float)
        if not np.isfinite(values).all() or not np.equal(values, np.floor(values)).all():
            raise ValueError(f"Sequence column {column!r} must contain finite integers")
        work[column] = values.astype(np.int64)

    output: list[dict[str, Any]] = []
    qc_rows: list[dict[str, int | float]] = []
    previous_segment: int | None = None
    previous_absolute: int | None = None
    previous_frame: int | None = None
    removed_reordered = 0
    inserted_missing = 0
    segment_transitions = 0
    for _, row_series in work.iterrows():
        row = row_series.to_dict()
        segment = int(row["segment_id"])
        absolute = int(row["absolute_seq_u64"])
        frame_seq = int(row["frame_seq"])
        if previous_segment is not None and segment != previous_segment:
            segment_transitions += 1
            previous_absolute = None
            previous_frame = None
        if previous_absolute is not None:
            delta = absolute - previous_absolute
            if delta <= 0:
                removed_reordered += 1
                continue
            gap = delta - 1
            for offset in range(1, delta):
                inserted = {column: np.nan for column in PYDISPLAY_SENSOR_COLUMNS}
                inserted["segment_id"] = segment
                inserted["absolute_seq_u64"] = previous_absolute + offset
                inserted["frame_seq"] = (int(previous_frame) + offset) % 65536
                inserted["parser_valid"] = False
                output.append(inserted)
                qc_rows.append(
                    {
                        "SampleIndex": len(output) - 1,
                        "Seq": previous_absolute + offset,
                        "ValidFlag": 0,
                        "InterpFlag": 1,
                        "GapLen": gap,
                        "MissingBefore": 0,
                    }
                )
                inserted_missing += 1
        else:
            gap = 0
        output.append(row)
        qc_rows.append(
            {
                "SampleIndex": len(output) - 1,
                "Seq": absolute,
                "ValidFlag": 1,
                "InterpFlag": 0,
                "GapLen": 0,
                "MissingBefore": gap,
            }
        )
        previous_segment = segment
        previous_absolute = absolute
        previous_frame = frame_seq
    frame = pd.DataFrame(output, columns=PYDISPLAY_SENSOR_COLUMNS)
    qc = pd.DataFrame(qc_rows)
    metadata = {
        "input_rows": input_rows,
        "output_rows": len(frame),
        "removed_empty_rows": removed_empty,
        "removed_parser_invalid_rows": removed_invalid,
        "removed_duplicate_or_reordered_rows": removed_reordered,
        "inserted_missing_rows": inserted_missing,
        "segment_transitions": segment_transitions,
    }
    return frame, qc, metadata


def _interpolate_finite(values: np.ndarray, name: str) -> np.ndarray:
    series = pd.Series(np.asarray(values, dtype=float)).replace([np.inf, -np.inf], np.nan)
    if not series.notna().any():
        raise ValueError(f"{name} has no finite values")
    result = series.interpolate(limit_direction="both").to_numpy(dtype=float)
    if not np.isfinite(result).all():
        raise ValueError(f"{name} contains unrecoverable values")
    return result


def _parser_valid_value(value: object) -> bool:
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, float, np.integer, np.floating)) and np.isfinite(value):
        return float(value) == 1.0
    return str(value).strip().lower() in {"true", "1"}


def _cell_sequence(payload: dict[str, Any], name: str, *, minimum: int) -> list[Any]:
    if name not in payload:
        raise ValueError(f"Calibration MAT missing field {name}")
    value = payload[name]
    items = list(value.ravel()) if isinstance(value, np.ndarray) else list(value) if isinstance(value, (list, tuple)) else [value]
    if len(items) < minimum:
        raise ValueError(f"Calibration field {name} must contain at least {minimum} cells")
    return items


def _mat_struct(value: Any, label: str) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise ValueError(f"{label} must be a MATLAB struct")
    return value


def _finite_scalar(entry: dict[str, Any], field: str, label: str) -> float:
    if field not in entry:
        raise ValueError(f"Calibration field missing: {label}")
    value = np.asarray(entry[field])
    if value.size != 1:
        raise ValueError(f"{label} must be scalar")
    scalar = float(value.reshape(-1)[0])
    if not np.isfinite(scalar):
        raise ValueError(f"{label} must be finite")
    return scalar


def _expect_integer(entry: dict[str, Any], field: str, expected: int, label: str) -> None:
    value = _finite_scalar(entry, field, label)
    if value != expected:
        raise ValueError(f"{label} must equal {expected}, got {value:g}")


def _has_data_row(path: Path) -> bool:
    try:
        with path.open("r", encoding="utf-8-sig", errors="replace") as handle:
            next(handle, None)
            return any(line.strip() for line in handle)
    except OSError:
        return False


def _issue(path: Path, category: str, reason: str) -> DiscoveryIssue:
    return DiscoveryIssue(path.name, path, category, reason)
